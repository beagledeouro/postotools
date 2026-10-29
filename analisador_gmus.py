#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
===============================================================================
 G-MUS Analyzer & Dashboard Generator - Atenção Primária à Saúde (APS / ACS)
 
 Módulo de inteligência demográfica, estratificação de risco clínico,
 auditoria cadastral (Previne Brasil) e geração de relatórios Excel multi-abas.
===============================================================================
"""

import os
import re
from datetime import datetime, date
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DIRETORIO_ATUAL = os.path.dirname(os.path.abspath(__file__))
ARQUIVO_PACIENTES = os.path.join(DIRETORIO_ATUAL, "Dados_Completos_Pacientes.xlsx")
ARQUIVO_DOMICILIOS = os.path.join(DIRETORIO_ATUAL, "Domicilios_Extraidos.xlsx")
ARQUIVO_RELATORIO_FINAL = os.path.join(DIRETORIO_ATUAL, "Relatorio_Microarea_Completo.xlsx")


# --- FUNÇÕES DE LIMPEZA E FORMATAÇÃO DE CAMPOS ---
def limpar_texto(v) -> str:
    if v is None or pd.isna(v):
        return ""
    s = str(v).strip()
    return "" if s.lower() in ["nan", "none", "null"] else s


def limpar_sexo(v) -> str:
    s = limpar_texto(v)
    if not s:
        return ""
    return re.sub(r"\s*\(.*", "", s).strip()


def limpar_rua(v) -> str:
    s = limpar_texto(v)
    if not s:
        return ""
    return re.sub(r"\s*\(\d+\)", "", s).strip()


def limpar_bairro(v) -> str:
    s = limpar_texto(v)
    if not s:
        return ""
    return re.sub(r"\s*\(\d+\)", "", s).strip()


def limpar_numero(v) -> str:
    s = limpar_texto(v)
    if not s:
        return "S/N"
    # Se for float (ex: 91.0)
    if s.endswith(".0"):
        s = s[:-2]
    return s.strip()


def limpar_documento(v) -> str:
    s = limpar_texto(v)
    if not s:
        return ""
    if s.endswith(".0"):
        s = s[:-2]
    return re.sub(r"\D", "", s)


def limpar_cpf(v) -> str:
    doc = limpar_documento(v)
    if not doc:
        return ""
    # Se o CPF perdeu zeros à esquerda no Excel/Pandas (ex: CPFs antigos ou que começam com 0)
    if 1 <= len(doc) <= 10:
        doc = doc.zfill(11)
    return doc


def formatar_cpf(v) -> str:
    doc = limpar_cpf(v)
    if len(doc) == 11:
        return f"{doc[:3]}.{doc[3:6]}.{doc[6:9]}-{doc[9:]}"
    return doc


def formatar_cns(v) -> str:
    doc = limpar_documento(v)
    if len(doc) == 15:
        return f"{doc[:3]} {doc[3:7]} {doc[7:11]} {doc[11:]}"
    return doc


def formatar_celular(ddd, num) -> str:
    d = limpar_documento(ddd)
    n = limpar_documento(num)
    if not n:
        return ""
    if len(n) == 9:
        n_fmt = f"{n[:5]}-{n[5:]}"
    elif len(n) == 8:
        n_fmt = f"{n[:4]}-{n[4:]}"
    else:
        n_fmt = n
    return f"({d}) {n_fmt}" if d else n_fmt


# --- CÁLCULOS DEMOGRÁFICOS E CLÍNICOS ---
def calcular_idade(dt_str) -> int:
    s = limpar_texto(dt_str)
    if not s:
        return None
    try:
        dt = datetime.strptime(s, "%d/%m/%Y").date()
        hoje = date.today()
        return hoje.year - dt.year - ((hoje.month, hoje.day) < (dt.month, dt.day))
    except Exception:
        return None


def classificar_faixa_etaria(idade) -> str:
    if idade is None:
        return "Não Informado"
    if idade <= 2:
        return "Bebê (0-2 anos)"
    elif idade <= 11:
        return "Criança (3-11 anos)"
    elif idade <= 17:
        return "Adolescente (12-17 anos)"
    elif idade <= 59:
        return "Adulto (18-59 anos)"
    else:
        return "Idoso (60+ anos)"


def classificar_risco_clinico(row: dict) -> tuple:
    """Calcula o risco clínico e lista as condições ativas do morador."""
    condicoes = []
    if str(row.get("Hipertensão", "")).strip().lower() == "sim":
        condicoes.append("Hipertensão")
    if str(row.get("Diabetico", "")).strip().lower() == "sim":
        condicoes.append("Diabetes")
    if str(row.get("Fumante", "")).strip().lower() == "sim":
        condicoes.append("Fumante")
    if str(row.get("Cardiaco", "")).strip().lower() == "sim":
        condicoes.append("Cardíaco")
    if str(row.get("Rins", "")).strip().lower() == "sim":
        condicoes.append("Rins")
    if str(row.get("Respiratorio", "")).strip().lower() == "sim":
        condicoes.append("Respiratório")
    if str(row.get("Câncer", "")).strip().lower() == "sim":
        condicoes.append("Câncer")

    qtd = len(condicoes)
    if qtd == 0:
        nivel = "Habitual"
        descricao = "Sem condição crônica declarada"
    elif qtd == 1:
        nivel = "Médio Risco"
        descricao = condicoes[0]
    elif qtd == 2:
        nivel = "Alto Risco"
        descricao = f"{condicoes[0]} + {condicoes[1]}"
    else:
        nivel = "Muito Alto Risco"
        descricao = f"{qtd} Condições ({', '.join(condicoes)})"

    return nivel, descricao


def checar_auditoria_cadastro(row: dict) -> str:
    """Detecta pendências no cadastro de acordo com as diretrizes da APS."""
    pendencias = []
    cpf = limpar_cpf(row.get("CPF", ""))
    if not cpf or len(cpf) != 11:
        pendencias.append("Sem CPF")

    cns = limpar_documento(row.get("CNS", ""))
    if not cns or len(cns) != 15:
        pendencias.append("Sem CNS")

    cel = limpar_documento(row.get("Celular", ""))
    if not cel or len(cel) < 8:
        pendencias.append("Sem Telefone")

    num = limpar_numero(row.get("Numero", ""))
    if not num or num == "S/N":
        pendencias.append("Sem Nº Logradouro")

    return ", ".join(pendencias) if pendencias else "Regular"


# --- ENRIQUECIMENTO DOS DADOS ---
def enriquecer_dataframe_pacientes(df: pd.DataFrame, df_dom: pd.DataFrame = None) -> pd.DataFrame:
    df_novo = df.copy()

    # 1. Limpeza de campos base
    df_novo["Sexo"] = df_novo["Sexo"].apply(limpar_sexo)
    df_novo["Rua"] = df_novo["Rua"].apply(limpar_rua)
    if "Bairro" in df_novo.columns:
        df_novo["Bairro"] = df_novo["Bairro"].apply(limpar_bairro)
    df_novo["Numero"] = df_novo["Numero"].apply(limpar_numero)
    df_novo["Complemento"] = df_novo["Complemento"].apply(limpar_texto)
    df_novo["Telefone_Formatado"] = df_novo.apply(lambda r: formatar_celular(r.get("DDD"), r.get("Celular")), axis=1)
    df_novo["CPF"] = df_novo["CPF"].apply(limpar_cpf)
    df_novo["CPF_Formatado"] = df_novo["CPF"].apply(formatar_cpf)
    df_novo["CNS_Formatado"] = df_novo["CNS"].apply(formatar_cns)

    # 2. Cálculos Demográficos e de Risco
    df_novo["Idade"] = df_novo["Nascimento"].apply(calcular_idade)
    df_novo["Faixa_Etaria"] = df_novo["Idade"].apply(classificar_faixa_etaria)

    riscos = df_novo.apply(classificar_risco_clinico, axis=1)
    df_novo["Nivel_Risco"] = [r[0] for r in riscos]
    df_novo["Detalhe_Risco"] = [r[1] for r in riscos]
    df_novo["Status_Cadastro"] = df_novo.apply(checar_auditoria_cadastro, axis=1)

    # 3. Mesclar dados de Domicílio / Visitas (se disponível)
    if df_dom is not None and not df_dom.empty and "Numero_Domicilio" in df_dom.columns:
        df_dom_map = df_dom.set_index("Numero_Domicilio").to_dict(orient="index")
        df_novo["Ultima_Visita_Domicilio"] = df_novo["Domicílio_ID"].astype(str).map(
            lambda x: df_dom_map.get(x, {}).get("Ultima_Visita", "")
        )
        df_novo["Dias_Atualizacao_Domicilio"] = df_novo["Domicílio_ID"].astype(str).map(
            lambda x: df_dom_map.get(x, {}).get("Dias_Atualizacao", "")
        )
        df_novo["Status_Domicilio"] = df_novo["Domicílio_ID"].astype(str).map(
            lambda x: df_dom_map.get(x, {}).get("Status", "Completo")
        )

    # 4. Reordenação inteligente das colunas
    colunas_finais = [
        "Domicílio_ID", "Código", "Nome", "Idade", "Faixa_Etaria", "Sexo", "Nascimento",
        "CPF_Formatado", "CNS_Formatado", "Telefone_Formatado",
        "Rua", "Numero", "Complemento",
        "Nivel_Risco", "Detalhe_Risco", "Status_Cadastro",
        "Hipertensão", "Diabetico", "Fumante", "Cardiaco", "Rins", "Respiratorio", "Câncer",
        "Mãe", "Pai"
    ]
    if "Ultima_Visita_Domicilio" in df_novo.columns:
        colunas_finais.extend(["Ultima_Visita_Domicilio", "Dias_Atualizacao_Domicilio", "Status_Domicilio"])

    colunas_presentes = [c for c in colunas_finais if c in df_novo.columns]
    outras = [c for c in df_novo.columns if c not in colunas_presentes and not c.endswith("_Formatado")]

    return df_novo[colunas_presentes + outras]


# --- ESTILIZAÇÃO E GERAÇÃO DO EXCEL MULTI-ABAS ---
def aplicar_estilo_tabela(ws, cor_cabecalho="1E3A8A"):
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"

    fill_cabecalho = PatternFill(start_color=cor_cabecalho, end_color=cor_cabecalho, fill_type="solid")
    font_cabecalho = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_corpo = Font(name="Calibri", size=10)
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_alerta = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    font_alerta = Font(name="Calibri", size=10, color="991B1B", bold=True)

    border_suave = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    # 1. Cabeçalho
    ws.row_dimensions[1].height = 28
    for cell in ws[1]:
        cell.fill = fill_cabecalho
        cell.font = font_cabecalho
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 2. Corpo
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        ws.row_dimensions[row_idx].height = 20
        eh_zebra = (row_idx % 2 == 0)

        for cell in row:
            cell.font = font_corpo
            cell.border = border_suave
            cell.alignment = Alignment(vertical="center")

            if eh_zebra:
                cell.fill = fill_zebra

            # Destaque para pendências ou risco alto
            val_str = str(cell.value or "")
            if "Sem CPF" in val_str or "Sem CNS" in val_str:
                cell.fill = fill_alerta
                cell.font = font_alerta
            elif "Muito Alto" in val_str or "Alto Risco" in val_str:
                cell.font = Font(name="Calibri", size=10, color="DC2626", bold=True)

    # 3. Auto-fit nas colunas
    for col in ws.columns:
        max_len = 0
        for cell in col:
            v = str(cell.value or "")
            if len(v) > max_len:
                max_len = len(v)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 11), 46)

    # Auto-filtro na linha de cabeçalho
    ws.auto_filter.ref = ws.dimensions


def criar_aba_dashboard(wb, df_pacientes: pd.DataFrame, df_domicilios: pd.DataFrame):
    ws = wb.create_sheet(title="📊 Dashboard & Indicadores", index=0)
    ws.views.sheetView[0].showGridLines = True

    # Paleta de Cores
    fill_titulo = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    font_titulo = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    fill_subtitulo = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    font_subtitulo = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_card = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    font_card_num = Font(name="Calibri", size=18, bold=True, color="1E3A8A")
    font_card_lbl = Font(name="Calibri", size=9, bold=True, color="475569")
    border_card = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    # Banner Principal
    ws.merge_cells("B2:I2")
    ws["B2"] = "PAINEL EPIDEMIOLÓGICO E GESTÃO DEMOGRÁFICA - MICROÁREA 4"
    ws["B2"].fill = fill_titulo
    ws["B2"].font = font_titulo
    ws["B2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 42

    # Cálculos dos KPIs
    tot_pacientes = len(df_pacientes)
    tot_domicilios = len(df_domicilios) if df_domicilios is not None else len(df_pacientes["Domicílio_ID"].unique())
    tot_hipertensos = len(df_pacientes[df_pacientes["Hipertensão"] == "Sim"])
    tot_diabeticos = len(df_pacientes[df_pacientes["Diabetico"] == "Sim"])
    tot_idosos = len(df_pacientes[df_pacientes["Idade"] >= 60])
    tot_criancas = len(df_pacientes[df_pacientes["Idade"] <= 11])
    tot_pendencias = len(df_pacientes[df_pacientes["Status_Cadastro"] != "Regular"])

    cards = [
        ("B", "C", "TOTAL MORADORES", str(tot_pacientes)),
        ("D", "E", "TOTAL DOMICÍLIOS", str(tot_domicilios)),
        ("F", "G", "HIPERTENSOS (HAS)", f"{tot_hipertensos} ({tot_hipertensos/tot_pacientes*100:.1f}%)" if tot_pacientes else "0"),
        ("H", "I", "DIABÉTICOS (DM)", f"{tot_diabeticos} ({tot_diabeticos/tot_pacientes*100:.1f}%)" if tot_pacientes else "0"),
    ]

    for c1, c2, label, valor in cards:
        ws.merge_cells(f"{c1}4:{c2}4")
        ws.merge_cells(f"{c1}5:{c2}5")

        ws[f"{c1}4"] = label
        ws[f"{c1}4"].fill = fill_subtitulo
        ws[f"{c1}4"].font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        ws[f"{c1}4"].alignment = Alignment(horizontal="center", vertical="center")

        ws[f"{c1}5"] = valor
        ws[f"{c1}5"].fill = fill_card
        ws[f"{c1}5"].font = font_card_num
        ws[f"{c1}5"].alignment = Alignment(horizontal="center", vertical="center")

        for col_letter in [c1, c2]:
            ws[f"{col_letter}4"].border = border_card
            ws[f"{col_letter}5"].border = border_card

    cards_2 = [
        ("B", "C", "IDOSOS (60+ ANOS)", f"{tot_idosos} ({tot_idosos/tot_pacientes*100:.1f}%)" if tot_pacientes else "0"),
        ("D", "E", "CRIANÇAS (0-11 ANOS)", f"{tot_criancas} ({tot_criancas/tot_pacientes*100:.1f}%)" if tot_pacientes else "0"),
        ("F", "G", "PENDÊNCIAS CADASTRO", f"{tot_pendencias} ({tot_pendencias/tot_pacientes*100:.1f}%)" if tot_pacientes else "0"),
        ("H", "I", "CADASTROS REGULARES", f"{tot_pacientes - tot_pendencias} ({(tot_pacientes - tot_pendencias)/tot_pacientes*100:.1f}%)" if tot_pacientes else "0"),
    ]

    for c1, c2, label, valor in cards_2:
        ws.merge_cells(f"{c1}7:{c2}7")
        ws.merge_cells(f"{c1}8:{c2}8")

        ws[f"{c1}7"] = label
        ws[f"{c1}7"].fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
        ws[f"{c1}7"].font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        ws[f"{c1}7"].alignment = Alignment(horizontal="center", vertical="center")

        ws[f"{c1}8"] = valor
        ws[f"{c1}8"].fill = fill_card
        ws[f"{c1}8"].font = font_card_num
        ws[f"{c1}8"].alignment = Alignment(horizontal="center", vertical="center")

        for col_letter in [c1, c2]:
            ws[f"{col_letter}7"].border = border_card
            ws[f"{col_letter}8"].border = border_card

    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 35
    ws.row_dimensions[7].height = 20
    ws.row_dimensions[8].height = 35

    # Tabela 1: Distribuição por Faixa Etária (Colunas B:D, a partir da linha 11)
    ws.merge_cells("B11:D11")
    ws["B11"] = "DISTRIBUIÇÃO POR FAIXA ETÁRIA"
    ws["B11"].fill = fill_titulo
    ws["B11"].font = font_subtitulo
    ws["B11"].alignment = Alignment(horizontal="center", vertical="center")

    ws["B12"] = "Faixa Etária"
    ws["C12"] = "População"
    ws["D12"] = "Percentual"
    for col_c in ["B", "C", "D"]:
        ws[f"{col_c}12"].font = Font(name="Calibri", size=10, bold=True)
        ws[f"{col_c}12"].fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        ws[f"{col_c}12"].alignment = Alignment(horizontal="center")

    faixas = df_pacientes["Faixa_Etaria"].value_counts()
    r_idx = 13
    for faixa, qtd in faixas.items():
        ws[f"B{r_idx}"] = str(faixa)
        ws[f"C{r_idx}"] = int(qtd)
        ws[f"D{r_idx}"] = f"{qtd / tot_pacientes * 100:.1f}%" if tot_pacientes else "0%"
        ws[f"B{r_idx}"].border = border_card
        ws[f"C{r_idx}"].border = border_card
        ws[f"D{r_idx}"].border = border_card
        ws[f"C{r_idx}"].alignment = Alignment(horizontal="center")
        ws[f"D{r_idx}"].alignment = Alignment(horizontal="center")
        r_idx += 1

    # Tabela 2: Estratificação de Risco Clínico (Colunas F:H, linha 11)
    ws.merge_cells("F11:H11")
    ws["F11"] = "ESTRATIFICAÇÃO DE RISCO CLÍNICO"
    ws["F11"].fill = fill_titulo
    ws["F11"].font = font_subtitulo
    ws["F11"].alignment = Alignment(horizontal="center", vertical="center")

    ws["F12"] = "Nível de Risco"
    ws["G12"] = "Pacientes"
    ws["H12"] = "Percentual"
    for col_c in ["F", "G", "H"]:
        ws[f"{col_c}12"].font = Font(name="Calibri", size=10, bold=True)
        ws[f"{col_c}12"].fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        ws[f"{col_c}12"].alignment = Alignment(horizontal="center")

    riscos = df_pacientes["Nivel_Risco"].value_counts()
    r_idx_risco = 13
    for risco, qtd in riscos.items():
        ws[f"F{r_idx_risco}"] = str(risco)
        ws[f"G{r_idx_risco}"] = int(qtd)
        ws[f"H{r_idx_risco}"] = f"{qtd / tot_pacientes * 100:.1f}%" if tot_pacientes else "0%"
        ws[f"F{r_idx_risco}"].border = border_card
        ws[f"G{r_idx_risco}"].border = border_card
        ws[f"H{r_idx_risco}"].border = border_card
        ws[f"G{r_idx_risco}"].alignment = Alignment(horizontal="center")
        ws[f"H{r_idx_risco}"].alignment = Alignment(horizontal="center")
        r_idx_risco += 1

    # Tabela 3: Principais Logradouros (Linha 20)
    linha_inicio_ruas = max(r_idx, r_idx_risco) + 2
    ws.merge_cells(f"B{linha_inicio_ruas}:E{linha_inicio_ruas}")
    ws[f"B{linha_inicio_ruas}"] = "TOP LOGRADOUROS DA MICROÁREA"
    ws[f"B{linha_inicio_ruas}"].fill = fill_titulo
    ws[f"B{linha_inicio_ruas}"].font = font_subtitulo
    ws[f"B{linha_inicio_ruas}"].alignment = Alignment(horizontal="center", vertical="center")

    ws[f"B{linha_inicio_ruas+1}"] = "Logradouro"
    ws[f"C{linha_inicio_ruas+1}"] = "Moradores"
    ws[f"D{linha_inicio_ruas+1}"] = "Hipertensos"
    ws[f"E{linha_inicio_ruas+1}"] = "Diabéticos"
    for col_c in ["B", "C", "D", "E"]:
        ws[f"{col_c}{linha_inicio_ruas+1}"].font = Font(name="Calibri", size=10, bold=True)
        ws[f"{col_c}{linha_inicio_ruas+1}"].fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        ws[f"{col_c}{linha_inicio_ruas+1}"].alignment = Alignment(horizontal="center")

    top_ruas = df_pacientes["Rua"].value_counts().head(8)
    r_rua = linha_inicio_ruas + 2
    for rua, qtd in top_ruas.items():
        sub_rua = df_pacientes[df_pacientes["Rua"] == rua]
        has_rua = len(sub_rua[sub_rua["Hipertensão"] == "Sim"])
        dm_rua = len(sub_rua[sub_rua["Diabetico"] == "Sim"])

        ws[f"B{r_rua}"] = str(rua)
        ws[f"C{r_rua}"] = int(qtd)
        ws[f"D{r_rua}"] = int(has_rua)
        ws[f"E{r_rua}"] = int(dm_rua)

        for col_c in ["B", "C", "D", "E"]:
            ws[f"{col_c}{r_rua}"].border = border_card
        ws[f"C{r_rua}"].alignment = Alignment(horizontal="center")
        ws[f"D{r_rua}"].alignment = Alignment(horizontal="center")
        ws[f"E{r_rua}"].alignment = Alignment(horizontal="center")
        r_rua += 1

    # Larguras das colunas do Dashboard
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 30
    ws.column_dimensions["G"].width = 16
    ws.column_dimensions["H"].width = 16
    ws.column_dimensions["I"].width = 18


# --- CONSTRUTOR GERAL DO RELATÓRIO EXCEL (PLANILHA ÚNICA CONSOLIDADA) ---
def gerar_relatorio_completo_excel(
    df_pacientes: pd.DataFrame = None,
    df_domicilios: pd.DataFrame = None,
    caminho_pacientes=None,
    caminho_domicilios=None,
    caminho_saida=ARQUIVO_RELATORIO_FINAL
) -> str:
    """Gera a pasta de trabalho Excel multi-abas estilizada e enriquecida em UMA ÚNICA PLANILHA."""
    print("\n" + "=" * 65)
    print("GERANDO PLANILHA ÚNICA CONSOLIDADA (RELATÓRIO MULTI-ABAS APS / ACS)")
    print("=" * 65)

    if df_pacientes is None:
        if caminho_pacientes and os.path.exists(caminho_pacientes):
            df_pacientes = pd.read_excel(
                caminho_pacientes,
                dtype={"CPF": str, "CNS": str, "Código": str, "Celular": str, "DDD": str, "Domicílio_ID": str, "Numero": str}
            )
        elif os.path.exists(caminho_saida):
            try:
                df_pacientes = pd.read_excel(
                    caminho_saida,
                    sheet_name="👥 Todos os Moradores",
                    dtype={"CPF": str, "CNS": str, "Código": str, "Celular": str, "DDD": str, "Domicílio_ID": str, "Numero": str}
                )
            except Exception:
                pass
        elif os.path.exists(ARQUIVO_PACIENTES):
            df_pacientes = pd.read_excel(
                ARQUIVO_PACIENTES,
                dtype={"CPF": str, "CNS": str, "Código": str, "Celular": str, "DDD": str, "Domicílio_ID": str, "Numero": str}
            )

    if df_pacientes is None or df_pacientes.empty:
        raise FileNotFoundError("Nenhum dado de pacientes encontrado para gerar a planilha única.")

    if df_domicilios is None:
        if caminho_domicilios and os.path.exists(caminho_domicilios):
            df_domicilios = pd.read_excel(caminho_domicilios, dtype=str)
        elif os.path.exists(caminho_saida):
            try:
                df_domicilios = pd.read_excel(caminho_saida, sheet_name="🏠 Domicílios & Visitas", dtype=str)
            except Exception:
                pass
        elif os.path.exists(ARQUIVO_DOMICILIOS):
            df_domicilios = pd.read_excel(ARQUIVO_DOMICILIOS, dtype=str)

    print(f"[+] Processando e enriquecendo {len(df_pacientes)} pacientes...")
    df_enriquecido = enriquecer_dataframe_pacientes(df_pacientes, df_domicilios)

    # Criação temporária via Pandas Writer
    temp_path = caminho_saida + ".tmp.xlsx"
    with pd.ExcelWriter(temp_path, engine="openpyxl") as writer:
        # Aba 1: Moradores Completo
        df_enriquecido.to_excel(writer, sheet_name="👥 Todos os Moradores", index=False)

        # Aba 2: Hipertensos e Diabéticos
        df_cronicos = df_enriquecido[(df_enriquecido["Hipertensão"] == "Sim") | (df_enriquecido["Diabetico"] == "Sim")].copy()
        df_cronicos.to_excel(writer, sheet_name="❤️ Hipertensos & Diabéticos", index=False)

        # Aba 3: Idosos (60+)
        df_idosos = df_enriquecido[df_enriquecido["Idade"] >= 60].sort_values(by=["Rua", "Numero"]).copy()
        df_idosos.to_excel(writer, sheet_name="👵 Idosos (60+ Anos)", index=False)

        # Aba 4: Crianças e Bebês (0 a 11 anos)
        df_criancas = df_enriquecido[df_enriquecido["Idade"] <= 11].sort_values(by="Idade").copy()
        df_criancas.to_excel(writer, sheet_name="👶 Bebês & Crianças (0-11)", index=False)

        # Aba 5: Roteiro de Visitas (Ordenado por Logradouro e Número para caminhada do ACS)
        df_roteiro = df_enriquecido.sort_values(by=["Rua", "Numero"]).copy()
        df_roteiro.to_excel(writer, sheet_name="🚶 Roteiro de Visitas (Por Rua)", index=False)

        # Aba 6: Auditoria de Pendências Cadastrais (Previne Brasil)
        df_pendencias = df_enriquecido[df_enriquecido["Status_Cadastro"] != "Regular"].copy()
        df_pendencias.to_excel(writer, sheet_name="⚠️ Pendências Cadastrais", index=False)

        # Aba 7: Domicílios e Visitas
        if df_domicilios is not None and not df_domicilios.empty:
            df_dom_clean = df_domicilios.copy()
            if "Logradouro" in df_dom_clean.columns:
                df_dom_clean["Logradouro"] = df_dom_clean["Logradouro"].apply(limpar_rua)
            df_dom_clean.to_excel(writer, sheet_name="🏠 Domicílios & Visitas", index=False)

    print("[+] Aplicando formatação visual profissional com openpyxl...")
    wb = openpyxl.load_workbook(temp_path)

    # 1. Aplica formatação de tabela nas abas de dados
    cores_abas = {
        "👥 Todos os Moradores": "1E3A8A",            # Azul Escuro
        "❤️ Hipertensos & Diabéticos": "991B1B",      # Vermelho Escuro
        "👵 Idosos (60+ Anos)": "15803D",             # Verde Escuro
        "👶 Bebês & Crianças (0-11)": "0369A1",       # Azul Petróleo
        "🚶 Roteiro de Visitas (Por Rua)": "4338CA",  # Índigo
        "⚠️ Pendências Cadastrais": "B45309",         # Âmbar / Laranja
        "🏠 Domicílios & Visitas": "0F766E",          # Verde Água
    }

    for s_name in wb.sheetnames:
        ws = wb[s_name]
        cor = cores_abas.get(s_name, "1E3A8A")
        aplicar_estilo_tabela(ws, cor_cabecalho=cor)

    # 2. Cria o Dashboard visual como primeira aba
    criar_aba_dashboard(wb, df_enriquecido, df_domicilios)

    wb.save(caminho_saida)
    if os.path.exists(temp_path):
        os.remove(temp_path)

    print(f"[✓] Relatório Multi-Abas concluído com sucesso!")
    print(f"[+] Arquivo disponível em: {caminho_saida}")
    return caminho_saida


# --- BUSCA RÁPIDA OFFLINE VIA TERMINAL ---
def buscar_offline(termo: str, arquivo=ARQUIVO_RELATORIO_FINAL):
    if not os.path.exists(arquivo):
        if os.path.exists(ARQUIVO_PACIENTES):
            arquivo = ARQUIVO_PACIENTES
        else:
            print(f"[!] Planilha única local não encontrada em: {arquivo}")
            return

    try:
        df = pd.read_excel(
            arquivo,
            sheet_name="👥 Todos os Moradores",
            dtype={"CPF": str, "CNS": str, "Código": str, "Celular": str, "DDD": str, "Domicílio_ID": str, "Numero": str}
        )
    except Exception:
        df = pd.read_excel(
            arquivo,
            dtype={"CPF": str, "CNS": str, "Código": str, "Celular": str, "DDD": str, "Domicílio_ID": str, "Numero": str}
        )
    termo_str = str(termo).strip()
    termo_num = re.sub(r"\D", "", termo_str)

    mask = (
        df["Nome"].astype(str).str.contains(termo_str, case=False, na=False) |
        df["Código"].astype(str).str.contains(termo_str, case=False, na=False) |
        df["Domicílio_ID"].astype(str).str.contains(termo_str, case=False, na=False) |
        df["Rua"].astype(str).str.contains(termo_str, case=False, na=False)
    )

    if termo_num:
        mask = mask | (
            df["CPF"].apply(limpar_cpf).str.contains(termo_num, na=False) |
            df["CNS"].astype(str).str.contains(termo_num, na=False)
        )

    res = df[mask]
    print("\n" + "=" * 65)
    print(f"RESULTADO DA BUSCA POR: '{termo_str}' ({len(res)} moradores encontrados)")
    print("=" * 65)

    if res.empty:
        print("[!] Nenhum morador encontrado com esse critério.")
        return

    for _, r in res.iterrows():
        idade = calcular_idade(r.get("Nascimento"))
        idade_str = f"{idade} anos" if idade is not None else "Idade N/D"
        sexo = limpar_sexo(r.get("Sexo"))
        nome = r.get("Nome", "")
        dom = r.get("Domicílio_ID", "")
        cpf = formatar_cpf(r.get("CPF", ""))
        cns = formatar_cns(r.get("CNS", ""))
        cel = formatar_celular(r.get("DDD"), r.get("Celular"))
        rua = limpar_rua(r.get("Rua", ""))
        num = limpar_numero(r.get("Numero", ""))
        comp = limpar_texto(r.get("Complemento", ""))
        comp_str = f" ({comp})" if comp else ""
        nivel_risco, det_risco = classificar_risco_clinico(r)
        status_cad = checar_auditoria_cadastro(r)

        print(f"\n👤 {nome} ({idade_str}, {sexo})")
        print(f"   🏠 Domicílio: {dom} | CPF: {cpf or 'NÃO INFORMADO'} | Celular: {cel or 'NÃO INFORMADO'}")
        print(f"   📍 Endereço: {rua}, Nº {num}{comp_str}")
        print(f"   🩺 Risco Clínico: {nivel_risco} ({det_risco})")
        print(f"   📋 Auditoria Cadastro: {status_cad}")
        print("   " + "-" * 60)


if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        if sys.argv[1] in ["--buscar", "-b"] and len(sys.argv) > 2:
            buscar_offline(" ".join(sys.argv[2:]))
        elif sys.argv[1] in ["--gerar", "-g"]:
            gerar_relatorio_completo_excel()
        else:
            print("Uso:")
            print("  python3 analisador_gmus.py --gerar               (Gera o Excel multi-abas)")
            print("  python3 analisador_gmus.py --buscar <nome/cpf>   (Busca morador offline)")
    else:
        gerar_relatorio_completo_excel()
